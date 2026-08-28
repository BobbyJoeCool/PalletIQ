"""Generate PalletIQ Manual E2E Test Checklist (.docx)."""

from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

doc = Document()

# -- Page setup: landscape for wider tables
for section in doc.sections:
    section.orientation = 1  # landscape
    section.page_width = Cm(29.7)
    section.page_height = Cm(21.0)
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.5)
    section.right_margin = Cm(1.5)

# -- Styles
style = doc.styles['Normal']
font = style.font
font.name = 'Calibri'
font.size = Pt(9)

# -- Title page
doc.add_paragraph()
doc.add_paragraph()
title = doc.add_paragraph('PalletIQ Floor App')
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
for run in title.runs:
    run.font.size = Pt(28)
    run.font.bold = True

subtitle = doc.add_paragraph('Manual E2E Test Checklist')
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
for run in subtitle.runs:
    run.font.size = Pt(18)

doc.add_paragraph()
legend = doc.add_paragraph()
legend.alignment = WD_ALIGN_PARAGRAPH.CENTER
legend.add_run('Status Key:  ').bold = True
legend.add_run('✓ Works    ✗ Doesn\'t Work    — Not Tested    ? Partial Works')
for run in legend.runs:
    run.font.size = Pt(12)

doc.add_paragraph()
info = doc.add_paragraph()
info.alignment = WD_ALIGN_PARAGRAPH.CENTER
info.add_run('Tester: ____________________    Date: ____________________    Version: ____________________')
for run in info.runs:
    run.font.size = Pt(11)

doc.add_page_break()

# -- Screen data: (screen_code, screen_name, test_items)
SCREENS = [
    ("LOGIN", "Login / Authentication", [
        "App opens to /login with no prior session restored",
        "ZnumPad: leading 'z' cannot be backspaced past",
        "OK button disabled until at least one char beyond 'z' prefix",
        "zNumber input capped at 7 total characters",
        "Nonexistent zNumber: error tone, 'zNumber not found', field resets to 'z'",
        "Valid zNumber: navigates to /pin, 'Welcome: {firstName}' displayed",
        "PIN: 4 digits auto-submit on 4th digit, no separate OK tap",
        "Incorrect PIN: error tone, 'Incorrect PIN', PIN clears, name stays",
        "Correct PIN: navigates to Home with replace:true (Back can't return to PIN)",
        "Direct nav to /pin with no route state: redirects to /login",
        "Back button on PIN screen: returns to /login with fresh entry",
        "Idle timeout fires after 15 minutes, logs out to /login",
        "Header Logout button: clears token/user, navigates to /login",
        "Wake Database button: countdown timer, 'Database ready' or error",
        "Reseed Test Data button: status text and completion summary",
        "Dev-tools styled amber, never touch the login message bar",
        "Field edit after error immediately clears the error message",
    ]),
    ("PIN", "PIN Entry", [
        "PIN screen shows worker's first name from login",
        "4-digit entry with auto-submit on 4th digit",
        "Incorrect PIN clears PIN field but keeps worker name",
        "Correct PIN navigates to Home (replace:true)",
        "Back button returns to Login with fresh zNumber entry",
        "Direct navigation to /pin without state redirects to /login",
    ]),
    ("HOME", "Home Screen Navigation", [
        "6-column grid renders: Production, GPM, Inventory, Location, Container, Reporting",
        "Welcome info message 'Signed in as {name}' on mount",
        "PIP button navigates to /pull",
        "SDP button navigates to /put/directed",
        "MNP button navigates to /put/manual",
        "ELA button navigates to /empty/aisle",
        "ELZ button navigates to /empty/zone",
        "STG button navigates to /stage",
        "PII button navigates to /pallet",
        "IID button navigates to /item",
        "PAR button navigates to /pallet/reinstate",
        "LII button navigates to /location",
        "WLH button navigates to /hold",
        "ISI button navigates to /storage-inquiry",
        "SAR button navigates to /staged-aisle",
        "IRP button navigates to /reporting/individual",
        "CII button navigates to /container",
        "PRQ button navigates to /reporting/pull-request",
        "Unbuilt function buttons (OCC, LRP, WTP) show info message, no nav",
        "Every button shows jump code badge and function name",
    ]),
    ("PIP", "Pallet ID Pull", [
        "Pull Function dropdown defaults to CA (Carton Air)",
        "Valid printed container scan: blue wash, pull data populates",
        "Bottom section renders with blank placeholders before any scan",
        "Mismatched pull function scan: 'Wrong function' error, red wash",
        "Nonexistent container scan: 'Label not found' error, red wash",
        "PULLED/CANCELED/PURGED container: 'Invalid status' error",
        "Pallet ID / UPC / Location fields disabled until CID is valid",
        "Correct Pallet ID verification: container marked PULLED, quantities deducted, history updated",
        "Mismatched Pallet ID: 'Incorrect Pallet ID', red wash",
        "Mismatched UPC: 'Invalid UPC', red wash",
        "Mismatched Location (aisle+bin): 'Invalid Location', all 3 boxes red",
        "Locked Aisle box rejects scan with different aisle segment",
        "Hold button with current+previous locations: picker popup with both options",
        "Change Pull Function mid-verification: 'Label not verified' warning, fields clear",
        "Demo footer: Valid Label, Label by Status, Invalid Label buttons work",
        "Success message persists through next label scan",
    ]),
    ("SDP", "System Directed Put", [
        "Valid Aisle entry: freight-type badges appear beneath Aisle field",
        "Nonexistent Aisle: error, field clears and refocuses, no badges",
        "Storage Code and Zone override fields hidden from Worker role",
        "Size override available to every role",
        "Valid Put Pending Pallet scan: SDPVerifyPutModal opens with directed location in red",
        "Pallet ID field stays disabled until Aisle is set",
        "Nonexistent Pallet ID: 'Pallet not found' error, red wash",
        "Pallet with no cartons: 'Invalid Pallet: No Cartons'",
        "Canceled pallet: 'Invalid Pallet: Canceled'",
        "Pallet blocked by pending pull: 'Invalid Pallet: Pull Pending'",
        "Aisle with no eligible locations: 'No eligible locations' error",
        "Already-stored pallet: 'directing as move' note",
        "Verify-Put Modal: correct confirm location scan = pallet stored, modal closes",
        "Verify-Put Modal: wrong confirm location = 'Wrong location' error",
        "Unassign button: reservation released, modal closes, Aisle/overrides kept, PID cleared",
        "Hold Location button: LockedHoldConfirmDialog opens with Hold Both / reason 70",
        "Hold Location confirm: reservation cleared, screen resets",
        "Hold Location cancel: return to Verify-Put Modal untouched",
        "Hand Put (XS): Exists Elsewhere button appears for IM+ with same-DPCI match",
        "Exists Elsewhere hidden from Worker role",
        "Exists Elsewhere pick: retarget (not immediate complete), Cancel/Return to Original appear",
        "15-second poll detects reservation expiry: modal closes with warning, full reset",
        "Back/Home/Jump/Logout disabled while reservation active (nav-locked)",
        "Consolidation timer: 15-minute expiry (not 5-minute)",
        "Consolidation: Cancel button labeled 'Unassign'",
        "Demo footer: Valid Pallet ID, Pallet ID by Status, Invalid Pallet ID buttons work",
    ]),
    ("MNP", "Manual Put", [
        "Single Scan Pallet ID field auto-focused on mount",
        "Valid pallet scan: pallet data displays, MNP_SCAN logged unconditionally",
        "Nonexistent pallet: 'Pallet not found', red wash, MNP_SCAN still logged",
        "Pallet with no cartons: error with red wash, value kept",
        "Already-stored pallet: info message 'proceeding as move'",
        "Valid destination Aisle+Bin: Level Modal opens",
        "Nonexistent destination: 'Location not found', boxes wash red",
        "Level Modal: enter level + confirm = put succeeds (no gates)",
        "Contraction gate: Worker hard-blocked; IM+ sees confirmation popup and can proceed",
        "Hold gate: HOLD_IN/HOLD_BOTH blocks Worker, IM+ can acknowledge; HOLD_PERM blocks all",
        "Occupied/Staged gate: same-DPCI offers Combine (IM+ only) or Cancel",
        "Occupied/Staged gate: different DPCI offers Proceed Anyway / Place Hold Both & Cancel / Cancel",
        "Consolidate (IM+, same-DPCI Combine): incoming zeroed, occupant quantities increased",
        "Cancel Put: return to ready state, MNP_CANCEL logged",
        "Abandonment via nav-away or idle-timeout fires MNP_CANCEL best-effort",
        "Hold quick-action for pallet's current location: HoldPanel opens, no default type/reason",
        "Demo footer: Pallet ID Demo Scanner (ready) and Location ID Demo Scanner (pallet_scanned)",
    ]),
    ("PII", "Pallet ID Info", [
        "Navigate to /pallet: Pallet ID field auto-focuses",
        "Navigate via ?id=<pid>: pallet loads immediately without auto-focus",
        "Valid pallet ID: read-only Loaded state with two-column layout",
        "Nonexistent pallet ID: 'Pallet not found', error tone, red wash",
        "'Go to Location ID' button present (disabled if no location); navigates to /location?id=",
        "Edit button visible only for IM+ roles; Worker never sees it",
        "Edit mode (IM+): fields seeded from loaded pallet with 'Current: {value}' indicators",
        "Save button enables only when a parsed value actually differs",
        "Save without Reason Code: 'A reason code is required' block",
        "Save with valid changes: success message, re-fetch, return to Loaded",
        "Expiration 1-3 months out: EXPIRATION_NEEDS_CONFIRM popup",
        "Expiration under 1 month: 'Expiration Date must be at least 1 month out' error",
        "VCP/SSP ratio validation: SSP must divide evenly into VCP",
        "SSPs on Pallet cap: must be less than one full carton's worth",
        "Cancel in Edit mode: discard changes, return to Loaded",
        "Scan new pallet while in Edit mode: unsaved changes discarded silently",
        "Session persistence: navigate away and back, last-loaded pallet restores read-only",
        "Demo footer: Valid Pallet ID, Pallet ID by Status, Invalid Pallet ID buttons work",
    ]),
    ("IID", "Item ID Lookup", [
        "Navigate to /item: Dept box auto-focuses",
        "Dept(3) -> Class(2) -> Item(4): auto-advance and auto-resolve",
        "UPC entry + confirm: item loads, DPCI boxes backfill",
        "DPCI entry: UPC field clears (asymmetric behavior)",
        "Nonexistent DPCI: 'Item not found', red wash on DPCI group",
        "Nonexistent UPC: 'Item not found', red wash on UPC",
        "Navigate via ?dpci= or ?upc=: pre-population and auto-resolve",
        "Read-only fields display: DPCI, UPC, Name, Short Desc, Desc, Retail, Cost, Weight, Storage, Conveyable",
        "'View Storage Locations' navigates to ISI with ?dpci=",
        "'Reinstate Pallet' button visible only for IM+ roles",
        "'Reinstate Pallet' (IM+): navigates to PAR with ?dpci=",
        "Demo footer: Valid DPCI, DPCI by Filter, Invalid DPCI; relabel to UPC when UPC focused",
    ]),
    ("PAR", "Pallet Reinstate", [
        "Navigate as Worker: 'Access Denied', no form, no demo buttons",
        "Navigate as IM+: full form renders with all fields visible",
        "Valid DPCI: Description populates, UPC clears",
        "Valid UPC: DPCI boxes backfill, Description populates",
        "Invalid DPCI: red wash on DPCI group, 'DPCI not found' error",
        "VCP/SSP with SSP not dividing VCP: group wash, 'SSP must divide evenly into VCP'",
        "Auto-advance chain: DPCI/UPC -> VCP -> SSP -> Size -> Cartons -> SSPs -> (Exp) -> Location",
        "Size is mandatory unless Location entered (then disabled/inherited)",
        "Single Pallet / Multiple Pallets mode toggle: appropriate fields appear",
        "Single Pallet: Location per-box validation (Aisle/Bin/Level existence)",
        "Occupied/on-hold/contracted location: amber inline flag + warn-then-allow at Create",
        "Storage Code mismatch (item vs location): warning, not block",
        "Expiration Month out of 1-12: individual box wash",
        "Too-soon Expiration Date: group wash and error",
        "Expiration Date required when item's requiresExpirationDate is true",
        "Create Pallet: confirm dialog with 5-line layout",
        "Confirm creation: success message, form clears, Reinstate Log entry added",
        "Multiple Pallets mode: Location disabled, N+1 pallet rows created",
        "Demo footer: Item Demo Scanner and Location Demo Scanner work",
    ]),
    ("LII", "Location ID Info", [
        "Navigate to /location: Aisle box auto-focuses",
        "Aisle(3) + Bin(3) + Level(2): auto-advance and location loads",
        "Full 8-digit barcode scan: immediate resolve regardless of focused box",
        "Navigate via ?id=: pre-population and auto-resolve",
        "Nonexistent location: 'Location not found', error tone, boxes clear and remount",
        "Detail displays: Location ID, Aisle, Bin, Level, Zone, Size, Storage Code, Status, Hold",
        "Pallet summary always renders with 'PALLET x/y' indicator (0/0 if unoccupied)",
        "Multi-occupant: Next/Prev buttons cycle through occupants",
        "CONTRACTED badge renders inline next to Status when contraction: true",
        "'Go to Pallet ID' navigates to /pallet?id=<pid>",
        "'Hold' navigates to /hold?id=<8-digit location id>",
        "Session persistence: navigate away and back, last-loaded location restores",
        "Demo footer: Valid Location, Location by Filter, Invalid Location buttons work",
    ]),
    ("WLH", "Warehouse Location Hold", [
        "Single Location / Range toggle visible only for IM+; Worker locked to Single",
        "Single Location mode: Location indicator + HoldPanel render with '---' placeholders",
        "Valid location entry: indicator and HoldPanel populate with live hold state",
        "Nonexistent location: 'Location not found', red group wash",
        "Hold Type 2x2 grid: button matching current hold type is disabled",
        "No Hold Type placeable while existing hold can't be removed by caller's role",
        "Select Hold Type + Reason Code + Confirm Hold: hold placed",
        "Different hold already exists: 'Replace existing hold?' confirmation dialog",
        "Remove Hold: hold removed, no reason code required",
        "Role gating: Worker = Hold Both only; Hold In/Out = IM+; Hold Perm = Lead+",
        "Hold Log records each action in both modes",
        "Range mode (IM+): Aisle, Start/End Bin, optionally Start/End Level",
        "Review Hold disabled when Start Bin > End Bin",
        "Review Hold: ConfirmDialog shows location count and range details",
        "Confirm Range Place: success message and per-bucket breakdown in Hold Log",
        "Range Release with Release Level gating (Perm clears all; Both clears In/Out; etc.)",
        "Per-bucket role gating on Range Release",
        "Reason Code uses shared ReasonCodeField with department/role prefix",
        "Session persistence: navigate away and back, last single-location id restores",
    ]),
    ("SAR", "Staged Aisle Report", [
        "Both columns load with 'Loading...' then populate",
        "'Most Staged' sorted descending by staged count",
        "'Staged Longest' sorted descending by age",
        "Each row shows Aisle id, freight-type badges, and metric",
        "Tap row to select: Directed Put and Stage Aisle buttons enable with aisle number",
        "Tap already-selected row: deselects, buttons disable",
        "Selection in one column reflects in the other for same aisle",
        "Directed Put button navigates to SDP with aisle pre-populated",
        "Stage Aisle button navigates to STG with aisle pre-populated",
        "Zero staged aisles: both columns show 'No staged locations in system'",
        "No auto-refresh; must reopen to see updated data",
        "No role gating; all roles see identical report",
    ]),
    ("ISI", "Item Storage Inquiry", [
        "Navigate to /storage-inquiry: Dept box auto-focuses",
        "Valid DPCI via Dept/Class/Item: item locations display sorted by aisle/bin/level",
        "Valid UPC: locations display, DPCI boxes backfill",
        "Nonexistent DPCI: 'Item not found', red wash on DPCI group",
        "Item found with zero locations: 'No locations currently storing this item'",
        "Short Description renders once above results",
        "Tap result row to select: 'Go to Location ID' and 'Go to Pallet ID' buttons appear",
        "Tap same row again: deselects, buttons hidden",
        "'Go to Location ID' navigates to /location?id=",
        "'Go to Pallet ID' navigates to /pallet?id=",
        "Navigate via ?dpci= or ?upc=: pre-population and auto-resolve",
        "Session persistence: navigate away and back, last search restores with selection",
        "Sticky header row stays visible while scrolling results",
        "Demo footer: Valid DPCI, DPCI by Filter, Invalid DPCI; relabel to UPC when UPC focused",
    ]),
    ("ELA", "Empty Locations by Aisle", [
        "No filter: idle prompt 'Enter a Storage Code...'",
        "Valid Storage Code: query auto-runs, 'Displaying {code}: {description}' banner",
        "Invalid Storage Code: 'Invalid Storage Code' error, red wash, no query",
        "Valid Size: results re-sort client-side by that size's empty count",
        "Invalid Size: 'Invalid Size' error, no query",
        "Size is a display/sort control, not a query filter (all aisles/sizes still shown)",
        "Zero-but-exists cells render as blue-washed 0(0), not blank",
        "Tap column header: sorting toggles direction; ascending pushes zeros to bottom",
        "Tap Size column header: that size value fills the Size field",
        "Tap row to select: View Zone Map and Stage Aisle buttons enable",
        "View Zone Map navigates to ELZ with aisle and storageCode",
        "Stage Aisle navigates to STG with aisle, storageCode, and size",
        "Zero-result query: 'No empty or staged locations found' message",
        "Session persistence: navigate away and back, filter and selection restore",
    ]),
    ("ELZ", "Empty Locations by Zone", [
        "No aisle: idle prompt 'Enter an Aisle to view the zone map'",
        "Valid Aisle: grid loads with all 8 zone/side columns (unfiltered by Storage Code)",
        "Nonexistent Aisle: 'Invalid Aisle' error with red wash",
        "Optional Storage Code: zone summary panel narrows to that code; grid stays unfiltered",
        "Invalid Storage Code: error message; code dropped but grid still loads",
        "Contracted cells render highlighted red in the grid",
        "Contracted locations excluded from zone summary panel counts",
        "Each column fills full height with entries weighted by Size",
        "Zone summary shows one column per Storage Code, badges sorted largest Size first",
        "Stage Aisle navigates to STG with aisle and storageCode",
        "Navigate from ELA's View Zone Map: aisle and storageCode pre-populate",
        "Session persistence: navigate away and back, last-viewed aisle restores",
    ]),
    ("STG", "Stage Aisle", [
        "Three stack boxes (On Deck, Next, Staging) and Master Control render",
        "Master Control Aisle/Storage Code/Size: all three stacks live-inherit values",
        "Storage Code validates against codes narrowed to current Aisle",
        "Arm per-field override toggle: field becomes editable, pre-filled, auto-focused",
        "Disarm override: field reverts to inherited display",
        "Staging Quantity filled: Locations panel fetches and shows destination bubbles",
        "Dynamic bubble sizing (1/2/3 columns based on count)",
        "Fewer locations than Quantity: red 'No Location' shortfall bubbles",
        "STAGE button: locations marked STAGED, log entry written, queue compacts forward",
        "Queue compaction: Next/On Deck slide into Staging; empty Staging inherits values",
        "Tap location bubble: reject/hold popup with default reason W70",
        "Confirm rejection: Hold Both placed, new suggestion fetched",
        "Unstage Aisle button visible only for IM+",
        "Unstage Aisle (IM+): modal lists freight types with toggles and quantity fields",
        "Apply Unstage/Restage: locations cleared and restaged per active rows",
        "Live Info Panel: ELZ format when Aisle present; ELA table when only Storage Code",
        "Tap row in ELA-format table: commits Aisle to Master Control",
        "Master Control Zone: staging begins in that zone",
        "Invalid Storage Code/Size/Aisle: value stays, red wash, error message",
        "Clear Forks clears all three stacks including Zone/zoneOverride",
        "Session persistence via StagingContext across navigation",
    ]),
    ("IRP", "Individual Reporting", [
        "All 9 production functions listed on open",
        "Bulk (BK) and Breakpack (BKP) permanently greyed and non-tappable",
        "Functions with activity sort by first-started time; inactive follow in fixed order",
        "Each active function row shows correct field set (counts, rates, hours, % to goal)",
        "Shift-wide time indicator: earliest assignment start through current time",
        "Tap active function row: navigates to hour-by-hour zoom-in view",
        "Zoom-in view: only assigned clock-hours listed (no blank rows)",
        "Back in zoom-in view: returns to summary",
        "Data scoped to logged-in worker only (no way to view another worker)",
        "No auto-refresh; re-entering triggers fresh pull",
        "API failure: 'Unable to load today\\'s data' fallback",
    ]),
    ("PRQ", "Pull Request", [
        "Navigate to /reporting/pull-request via Home button or jump code PRQ",
        "Empty state: prompt displays until Aisle + Status are filled",
        "Aisle field: 3-digit numpad entry, existence check via useAisleField",
        "End Aisle field: validates higher than Start Aisle, shows error if not",
        "Workstation field: entering clears Aisle fields (mutual exclusion)",
        "Aisle entry clears Workstation field (mutual exclusion)",
        "Status dropdown: only Available, Printed, Verified options",
        "Live filtering: results populate as soon as Aisle + Status are valid",
        "L1 summary: rows grouped by Batch Date × Pull Function",
        "L1 sort: Batch Date ascending, then function order CA → CF → FP → BK",
        "L1 CA/CF/FP row: Cartons, Locations, Density fields display",
        "L1 BK row: Full Pallets, Pallet Density, Loose Cartons, Carton Density, Locations",
        "Select-then-zoom-in: tap highlights row, Zoom In button or double-tap drills",
        "L1 → L2 (range/workstation): per-aisle breakdown shows for selected batch/function",
        "L1 → L3 (single aisle): skips L2, goes directly to location-level",
        "L2 per-aisle: same field format as L1, scoped to one aisle",
        "L2 → L3: zoom into selected aisle row",
        "L3 location-level: rows show Aisle-Bin-Level with Cartons/Pallets/SSPs",
        "L3 additional filters: Bin Start, Bin End, Level",
        "L3 sort: Aisle, Bin, Level ascending",
        "Back button: returns to previous level, clears selection",
        "Zoom In button: disabled when no row selected",
        "Summary line updates at every level showing current scope",
        "Pull detail popup: opens on L3 zoom-in",
        "Popup header: tappable Location → LII, DPCI → IID, Pallet ID → PII",
        "Popup pull metadata: pull type, batch date, pull quantity",
        "Popup container list: each CID with status badge and quantity",
        "Tap container CID: immediate navigation to CII",
        "Back from CII returns to PRQ popup",
        "Session persistence: filters/level survive navigate-away and back",
        "No data: 'No pull data found' message at each level",
    ]),
    ("CII", "Container ID Inquiry", [
        "Navigate to /container: Container ID field renders, no persistent numpad",
        "Numpad opens on field focus, dismisses on blur/lookup",
        "Valid Full Case (91) CID: detail loads with pull info, quantities, pallet link",
        "Valid Overpack (92) CID: detail loads with open/closed state, contents list",
        "Valid SSP Pull Master (93) CID: detail loads with children list, cascade cancel info",
        "Valid SSP Unit (94) CID: read-only detail with source master link",
        "Valid Stray Each (95) CID: detail loads with destination, found-by info",
        "Compressed CID input expanded correctly to canonical 35-digit form",
        "Nonexistent CID: error tone, 'Container not found', red wash on field",
        "Type label and status badge render in header next to CID field",
        "Event log (right column) populates with timestamped entries",
        "Event log shows reason codes, user names, child CID links, store/qty changes",
        "Tap DPCI link: navigates to IID with dpci param",
        "Tap Pallet ID link: navigates to PII with pid param",
        "Tap Location link: navigates to LII with loc param",
        "Tap child/source CID link: reloads CII with that CID",
        "Cancel button (Full Case, IM+): visible only for AVAILABLE/VERIFIED status",
        "Cancel modal: ReasonCodeField (CONTAINER_CANCEL domain), optional note, confirm flow",
        "Cancel SSP Master (IM+): cascade warning with PENDING child count",
        "Force Close overpack (IM+): visible when open and not purged",
        "Reopen overpack (IM+): visible when closed and not purged",
        "Overpack action modal: ReasonCodeField (OVERPACK_STATUS domain), confirm flow",
        "Reassign stray each (IM+): visible when status is PENDING",
        "Reassign modal: store selector excludes current store, STRAY_REASSIGN domain",
        "All action buttons hidden from Worker role",
        "Success messages display via MessageBar after actions",
        "Error messages display via MessageBar on action failure",
        "Session persistence: navigate away and back, last-loaded container restores",
        "Demo scanner bar provides sample CIDs for all five container types",
        "Re-scan/re-type CID while loaded: loads new container without nav-away",
    ]),
    ("LOG", "Activity Log Overlay", [
        "Tap 'Activity' in Header on any authenticated screen: overlay opens",
        "Overlay shows 'Loading...' then populates newest-first",
        "Entries scoped to logged-in worker's own zNumber only",
        "12-hour rolling window from stored timestamps",
        "Hidden action types (RESERVE, MNP_SCAN, RES_TMOUT, STAGE) filtered out",
        "Each entry: bold red function tag, timestamp, color-coded detail lines",
        "Tap Pallet ID / Location ID / DPCI token: navigates to PII/LII/IID, overlay closes",
        "Tap Close or dimmed backdrop: overlay dismisses, underlying screen unchanged",
        "Zero entries: 'No activity in the last 12 hours'",
        "API failure: 'Couldn\\'t load activity -- please try again'",
        "Activity button disabled when underlying screen has nav lock (e.g., SDP reservation)",
    ]),
    ("SHARED", "Shared Chrome / Cross-Screen", [
        "Header renders on all authenticated screens with correct screen title",
        "Back button navigates to previous screen",
        "Home button navigates to / (Home)",
        "Jump code entry + confirm navigates to correct screen",
        "Logout button clears session and navigates to /login",
        "Idle timeout (15 min) fires on all screens",
        "Numpad panel opens/closes correctly on field focus/blur",
        "Keyboard panel opens/closes correctly on text fields",
        "ScaleToFit: app scales to fit viewport at 1366x1024",
        "Demo Scanner Bar: all 3 buttons (Valid/by-Filter/Invalid) work per screen",
        "Error tone plays on all validation errors across screens",
        "Success tone plays on successful operations",
        "Role gating consistent across all screens (Worker vs IM vs Lead)",
    ]),
]


def set_cell_shading(cell, color_hex):
    """Set cell background color."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shading = tcPr.makeelement(qn('w:shd'), {
        qn('w:fill'): color_hex,
        qn('w:val'): 'clear',
    })
    tcPr.append(shading)


def add_screen_page(screen_code, screen_name, items):
    # Header
    heading = doc.add_heading(f'{screen_code} — {screen_name}', level=1)
    heading.runs[0].font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)

    # Info line
    info = doc.add_paragraph()
    info.add_run('Tester: ____________    Date: ____________    Version: ____________').font.size = Pt(9)

    # Table: # | Test Item | Status | Comments
    table = doc.add_table(rows=1, cols=4)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    # Set column widths
    widths = [Cm(0.8), Cm(16), Cm(2.5), Cm(6)]
    for i, width in enumerate(widths):
        table.columns[i].width = width

    # Header row
    header_cells = table.rows[0].cells
    headers = ['#', 'Test Item', 'Status', 'Comments']
    for i, (cell, text) in enumerate(zip(header_cells, headers)):
        cell.text = text
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.font.bold = True
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        set_cell_shading(cell, '1A1A2E')

    # Data rows
    for idx, item in enumerate(items, 1):
        row = table.add_row()
        cells = row.cells

        # Number
        cells[0].text = str(idx)
        cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

        # Test item
        cells[1].text = item

        # Status (empty, to be filled: checkmark/x/dash/?)
        cells[2].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

        # Comments (empty)
        # leave blank

        # Style all cells
        for cell in cells:
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9)

        # Alternate row shading
        if idx % 2 == 0:
            for cell in cells:
                set_cell_shading(cell, 'F0F0F5')


# Generate all screen pages
for i, (code, name, items) in enumerate(SCREENS):
    add_screen_page(code, name, items)
    if i < len(SCREENS) - 1:
        doc.add_page_break()

# Save
output_path = '/Users/Bob/csd/Side-Projects/PalletIQ/Documentation/PalletIQ-Manual-E2E-Checklist.docx'
doc.save(output_path)
print(f'Saved to {output_path}')
